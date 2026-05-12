from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.util import Pt


def parse_outline(path: str) -> list[tuple[str, list[str]]]:
    text = Path(path).read_text(encoding="utf-8")
    slides: list[tuple[str, list[str]]] = []
    current_title: str | None = None
    bullets: list[str] = []

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("## Slide"):
            if current_title:
                slides.append((current_title, bullets))
            current_title = stripped.replace("##", "").strip()
            bullets = []
        elif stripped.startswith("- "):
            bullets.append(stripped[2:].strip())

    if current_title:
        slides.append((current_title, bullets))

    return slides


def export_pptx(outline_path: str, output_path: str) -> str:
    prs = Presentation()
    slides = parse_outline(outline_path)

    if not slides:
        slides = [("PMFlow Codex", ["No slides found in outline"])]

    for title, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title

        body = slide.placeholders[1]
        text_frame = body.text_frame
        text_frame.clear()

        for i, bullet in enumerate(bullets or ["TBD"]):
            p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = Pt(22)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return str(output)
